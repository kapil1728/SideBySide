"""End-to-end tests using generated PDF and DOCX sample documents."""

from io import BytesIO
from zipfile import ZipFile

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from app import app

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_homepage_uses_sidebyside_branding() -> None:
    app.config.update(TESTING=True)
    response = app.test_client().get("/")

    assert response.status_code == 200
    assert b"SideBySide | Document comparison" in response.data
    assert b"SideBySide home" in response.data


def test_favicon_is_served() -> None:
    response = app.test_client().get("/static/favicon.svg")

    assert response.status_code == 200
    assert b"<svg" in response.data


def make_docx(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_pdf(text: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def test_post_compares_docx_and_pdf_and_returns_highlighted_downloads() -> None:
    app.config.update(TESTING=True)
    client = app.test_client()
    text1 = "A bright red balloon floated over the quiet river."
    text2 = "The bright red balloon floated over the quiet river today."
    response = client.post(
        "/",
        data={
            "word_count": "5",
            "file1": (BytesIO(make_docx(text1)), "first.docx", DOCX_MIME),
            "file2": (BytesIO(make_pdf(text2)), "second.pdf", PDF_MIME),
        },
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    page = response.get_data(as_text=True)
    assert "matched passage" in page
    assert "bright red balloon floated over the quiet river" in page
    assert "<mark>bright</mark>" in page
    assert "PDF page 1" in page
    assert "File 1 highlighted" in page
    assert "File 2 highlighted" in page

    with client.get("/download/" + _request_id(page) + "/1") as download:
        assert download.status_code == 200
        assert "first_highlighted.docx" in download.headers["Content-Disposition"]
        highlighted_docx = Document(BytesIO(download.data))
        highlighted_runs = [
            run for paragraph in highlighted_docx.paragraphs for run in paragraph.runs
            if run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        ]
        assert highlighted_runs

    pdf_link = _download_url(page, 2)
    highlighted_pdf = client.get(pdf_link)
    assert highlighted_pdf.status_code == 200
    assert "second_highlighted.pdf" in highlighted_pdf.headers["Content-Disposition"]
    pdf = fitz.open(stream=highlighted_pdf.data, filetype="pdf")
    assert list(pdf[0].annots())
    pdf.close()


def test_empty_pdf_returns_a_clear_error() -> None:
    app.config.update(TESTING=True)
    client = app.test_client()
    empty_pdf = fitz.open()
    empty_pdf.new_page()
    empty_data = empty_pdf.tobytes()
    empty_pdf.close()
    response = client.post(
        "/",
        data={
            "word_count": "2",
            "file1": (BytesIO(empty_data), "empty.pdf", PDF_MIME),
            "file2": (BytesIO(make_docx("some readable words here")), "text.docx", DOCX_MIME),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 400
    assert b"empty or scanned PDF" in response.data


def _request_id(page: str) -> str:
    return _download_url(page, 1).split("/")[-2]


def _download_url(page: str, number: int) -> str:
    import re

    match = re.search(r'href="(/download/[0-9a-f]{32}/' + str(number) + r')"', page)
    assert match
    return match.group(1)
